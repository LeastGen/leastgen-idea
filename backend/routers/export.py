"""IdeaFlow — export endpoints.

PDF and DOCX export for idea cards.
"""

from __future__ import annotations

import io
import os
import re
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel

from backend.database import get_user_by_id

router = APIRouter()

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
RUN_DIR = PROJECT_ROOT / "ideaspark_run"


# ── Helpers ──


def _get_card_path(run_id: str) -> Path | None:
    """Get the path to an idea card markdown file."""
    card = RUN_DIR / run_id / "phase4" / "idea.std.en.md"
    if card.exists():
        return card
    return None


def _md_to_text(md: str) -> str:
    """Convert markdown to plain text."""
    # Remove code fences
    text = re.sub(r"```[\s\S]*?```", "", md)
    # Remove images
    text = re.sub(r"!\[.*?\]\(.*?\)", "", text)
    # Remove links, keep text
    text = re.sub(r"\[([^\]]*)\]\(.*?\)", r"\1", text)
    # Remove HTML tags
    text = re.sub(r"<[^>]+>", "", text)
    # Remove KaTeX math
    text = re.sub(r"\$\$[\s\S]*?\$\$", "[equation]", text)
    text = re.sub(r"\$[^$]*?\$", "[inline eq]", text)
    # Remove markdown headers
    text = re.sub(r"^#{1,6}\s+", "", text, flags=re.MULTILINE)
    # Remove bold/italic
    text = re.sub(r"\*{1,3}(.*?)\*{1,3}", r"\1", text)
    return text.strip()


def _md_to_html(md: str) -> str:
    """Convert markdown to simple HTML."""
    html = md
    # Code blocks
    html = re.sub(r"```(\w*)\n([\s\S]*?)```", r"<pre><code>\2</code></pre>", html)
    # Headers
    html = re.sub(r"^### (.+)$", r"<h3>\1</h3>", html, flags=re.MULTILINE)
    html = re.sub(r"^## (.+)$", r"<h2>\1</h2>", html, flags=re.MULTILINE)
    html = re.sub(r"^# (.+)$", r"<h1>\1</h1>", html, flags=re.MULTILINE)
    # Bold
    html = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", html)
    # Italic
    html = re.sub(r"\*(.+?)\*", r"<em>\1</em>", html)
    # Inline code
    html = re.sub(r"`([^`]+)`", r"<code>\1</code>", html)
    # Links
    html = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r'<a href="\2">\1</a>', html)
    # Horizontal rules
    html = re.sub(r"^---$", r"<hr>", html, flags=re.MULTILINE)
    # Paragraphs (double newlines)
    html = re.sub(r"\n\n", r"</p><p>", html)
    html = f"<p>{html}</p>"
    return html


# ── Routes ──


@router.get("/export/{run_id}/pdf")
async def export_pdf(run_id: str):
    """Export an idea card as PDF."""
    card_path = _get_card_path(run_id)
    if not card_path:
        raise HTTPException(status_code=404, detail="Idea card not found")

    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm, mm
        from reportlab.platypus import (
            ListFlowable,
            ListItem,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
        )
    except ImportError:
        raise HTTPException(status_code=501, detail="PDF export not available (reportlab not installed)")

    md = card_path.read_text(encoding="utf-8")
    text = md  # We'll use the markdown directly

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title2", parent=styles["Title"], fontSize=18, spaceAfter=12)
    h1_style = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=16, spaceAfter=8, spaceBefore=16)
    h2_style = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=14, spaceAfter=6, spaceBefore=12)
    body_style = ParagraphStyle("Body2", parent=styles["Normal"], fontSize=10, leading=14, spaceAfter=6)

    story = []
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            story.append(Spacer(1, 4))
        elif line.startswith("# "):
            story.append(Paragraph(line[2:], title_style))
        elif line.startswith("## "):
            story.append(Paragraph(line[3:], h1_style))
        elif line.startswith("### "):
            story.append(Paragraph(line[4:], h2_style))
        elif line.startswith("- "):
            story.append(Paragraph(f"• {line[2:]}", body_style))
        else:
            story.append(Paragraph(line, body_style))

    doc.build(story)
    pdf_bytes = buf.getvalue()
    buf.close()

    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="idea-card-{run_id}.pdf"'},
    )


@router.get("/export/{run_id}/docx")
async def export_docx(run_id: str):
    """Export an idea card as DOCX."""
    card_path = _get_card_path(run_id)
    if not card_path:
        raise HTTPException(status_code=404, detail="Idea card not found")

    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
    except ImportError:
        raise HTTPException(status_code=501, detail="DOCX export not available (python-docx not installed)")

    md = card_path.read_text(encoding="utf-8")

    doc = Document()

    # Set default font — pyright: ignore[reportAttributeAccessIssue]
    style = doc.styles["Normal"]
    style.font.name = "Calibri"  # pyright: ignore
    style.font.size = Pt(11)  # pyright: ignore

    for line in md.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        else:
            doc.add_paragraph(line)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="idea-card-{run_id}.docx"'},
    )


@router.get("/export/{run_id}/md")
async def export_md(run_id: str):
    """Export an idea card as raw markdown."""
    card_path = _get_card_path(run_id)
    if not card_path:
        raise HTTPException(status_code=404, detail="Idea card not found")

    md = card_path.read_text(encoding="utf-8")
    return Response(
        content=md,
        media_type="text/markdown",
        headers={"Content-Disposition": f'attachment; filename="idea-card-{run_id}.md"'},
    )